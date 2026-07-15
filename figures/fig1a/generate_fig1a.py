#!/usr/bin/env python3
"""
Generate Figure 1A: PGC induction experimental schematic (minimalist B&W).

Outputs (all under this folder):
  - fig1a_pgc_induction.pptx  — native PowerPoint shapes (best for editing)
  - fig1a_pgc_induction.svg   — vector (Illustrator / PPT Insert > Picture)
  - fig1a_pgc_induction.png   — raster preview

Re-run after tweaking the geometry constants below:
  python3 generate_fig1a.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

OUT_DIR = Path(__file__).resolve().parent
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x33, 0x33, 0x33)

SLIDE_W = 7.4
SLIDE_H = 4.1

Y_MAIN = 1.35
LINE_THIN = Pt(1.15)
LINE_MED = Pt(1.5)
DOT_R = 0.052
X_MARK = 0.085

# Main-timeline x positions (inches)
X0 = 0.70
X_CA_END = 1.55
X_BMP = 2.25
X_2S = 2.95
X_INDUC_END = 6.85

# Four sampling dots under BMP+SCF+EGF (after 2S)
DOT_SPACING = 0.90
DOTS_MAIN = [3.55, 3.55 + DOT_SPACING, 3.55 + 2 * DOT_SPACING, 3.55 + 3 * DOT_SPACING]

BRANCH_YS = [2.20, 2.85, 3.45]
# Slightly shorter than inter-dot spacing so terminal • are not
# pierced by the next timepoint's dashed drop-line.
BRANCH_LEN = 0.70


def _dash(shape, dash="dash"):
    ln = shape.line._ln
    for child in list(ln):
        if child.tag.endswith("}prstDash"):
            ln.remove(child)
    ln.append(
        parse_xml(
            f'<a:prstDash xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            f'val="{dash}"/>'
        )
    )


def _fill_black(shape):
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLACK
    shape.line.fill.background()


def add_text(slide, left, top, width, height, text, size=10, bold=False,
             align=PP_ALIGN.CENTER, color=BLACK, font_name="Arial"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return box


def add_hline(slide, x1, x2, y, weight=LINE_THIN, dash=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y), Inches(x2), Inches(y)
    )
    shape.line.color.rgb = BLACK
    shape.line.width = weight
    if dash:
        _dash(shape, dash)
    return shape


def add_vline(slide, x, y1, y2, weight=LINE_THIN, dash=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y1), Inches(x), Inches(y2)
    )
    shape.line.color.rgb = BLACK
    shape.line.width = weight
    if dash:
        _dash(shape, dash)
    return shape


def add_tick(slide, x, y, half=0.075, weight=LINE_THIN):
    return add_vline(slide, x, y - half, y + half, weight=weight)


def add_dot(slide, x, y, r=DOT_R):
    oval = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x - r), Inches(y - r), Inches(2 * r), Inches(2 * r)
    )
    _fill_black(oval)
    return oval


def add_x_mark(slide, x, y, size=X_MARK, weight=LINE_MED):
    s = size / 2
    for (x1, y1, x2, y2) in (
        (x - s, y - s, x + s, y + s),
        (x - s, y + s, x + s, y - s),
    ):
        c = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
        )
        c.line.color.rgb = BLACK
        c.line.width = weight


def add_bracket(slide, x1, x2, y, tip=0.10):
    add_vline(slide, x1, y, y + tip)
    add_hline(slide, x1, x2, y + tip)
    add_vline(slide, x2, y, y + tip)


def add_down_arrow(slide, x, y1, y2, head=0.085):
    add_vline(slide, x, y1, y2 - head * 0.35, weight=LINE_THIN)
    tri = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(x - head * 0.55),
        Inches(y2 - head),
        Inches(head * 1.1),
        Inches(head),
    )
    tri.rotation = 180
    _fill_black(tri)


def build_pptx(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Panel letter
    add_text(slide, 0.15, 0.10, 0.32, 0.30, "A", size=14, bold=True, align=PP_ALIGN.LEFT)

    # Title
    add_text(
        slide, 0.42, 0.28, 2.0, 0.28,
        "PGC induction:", size=12, bold=True, align=PP_ALIGN.LEFT,
    )

    # "2d FA" annotates the title word "induction" (not a timeline event)
    add_text(slide, 1.05, 0.02, 0.70, 0.20, "2d FA", size=9, align=PP_ALIGN.CENTER)
    add_down_arrow(slide, 1.40, 0.20, 0.42, head=0.07)

    # --- Main timeline ---
    add_hline(slide, X0, X_INDUC_END, Y_MAIN, weight=LINE_MED)

    # CA ticks + bracket
    add_tick(slide, X0, Y_MAIN, half=0.08)
    add_tick(slide, X_CA_END, Y_MAIN, half=0.08)
    add_bracket(slide, X0, X_CA_END, Y_MAIN + 0.14, tip=0.10)
    add_text(slide, X0 - 0.05, Y_MAIN + 0.26, X_CA_END - X0 + 0.10, 0.22, "(CA)", size=9)

    # BMP
    add_dot(slide, X_BMP, Y_MAIN)
    add_text(slide, X_BMP - 0.35, Y_MAIN - 0.34, 0.70, 0.24, "BMP", size=10, bold=True)

    # 2S
    add_tick(slide, X_2S, Y_MAIN, half=0.11)
    add_text(slide, X_2S - 0.25, Y_MAIN + 0.12, 0.50, 0.22, "2S", size=9)

    # Induction cocktail label
    add_text(
        slide, X_2S + 0.15, Y_MAIN - 0.40, X_INDUC_END - X_2S - 0.15, 0.24,
        "BMP + SCF + EGF", size=10, bold=True,
    )

    for xd in DOTS_MAIN:
        add_dot(slide, xd, Y_MAIN)

    # Diverted germ-layer conditions from first three sampling points
    for xd, yb in zip(DOTS_MAIN[:3], BRANCH_YS):
        add_vline(slide, xd, Y_MAIN + DOT_R + 0.02, yb, weight=LINE_THIN, dash="dash")
        add_hline(slide, xd, xd + BRANCH_LEN, yb, weight=LINE_MED)
        add_x_mark(slide, xd, yb)
        add_dot(slide, xd + BRANCH_LEN, yb)

    # Legend
    lx, ly = 0.55, 3.65
    add_x_mark(slide, lx, ly, size=0.075)
    add_text(
        slide, lx + 0.16, ly - 0.10, 5.0, 0.22,
        "Ecto (SB + LDN)  ·  Meso  ·  Endo",
        size=8, align=PP_ALIGN.LEFT, color=GRAY,
    )
    add_dot(slide, lx, ly + 0.28, r=0.042)
    add_text(
        slide, lx + 0.16, ly + 0.18, 2.2, 0.22,
        "fix & stain",
        size=8, align=PP_ALIGN.LEFT, color=GRAY,
    )

    prs.save(str(path))
    print(f"Wrote {path}")


def build_svg(path: Path) -> None:
    W, H = int(SLIDE_W * 100), int(SLIDE_H * 100)

    def ix(v):
        return v * 100

    def iy(v):
        return v * 100

    y = iy(Y_MAIN)
    r = DOT_R * 100
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',
        '<rect width="100%" height="100%" fill="white"/>',
        "<style>"
        "text{font-family:Arial,Helvetica,sans-serif;fill:#000}"
        ".b{font-size:14px;font-weight:700}"
        ".h{font-size:12px;font-weight:700}"
        ".t{font-size:11px}"
        ".s{font-size:10px;fill:#333}"
        "line{stroke:#000;stroke-linecap:square;fill:none}"
        ".thin{stroke-width:1.15}.med{stroke-width:1.55}"
        ".dash{stroke-dasharray:4 3}"
        "</style>",
        f'<text x="{ix(0.15)}" y="{iy(0.30)}" class="b">A</text>',
        f'<text x="{ix(0.42)}" y="{iy(0.46)}" class="b">PGC induction:</text>',
        f'<text x="{ix(1.40)}" y="{iy(0.16)}" class="t" text-anchor="middle">2d FA</text>',
        f'<line class="thin" x1="{ix(1.40)}" y1="{iy(0.20)}" x2="{ix(1.40)}" y2="{iy(0.35)}"/>',
        f'<polygon points="{ix(1.40)},{iy(0.42)} {ix(1.40)-4.5},{iy(0.34)} {ix(1.40)+4.5},{iy(0.34)}" fill="#000"/>',
        # main line
        f'<line class="med" x1="{ix(X0)}" y1="{y}" x2="{ix(X_INDUC_END)}" y2="{y}"/>',
    ]

    for xt in (X0, X_CA_END):
        parts.append(f'<line class="thin" x1="{ix(xt)}" y1="{y-8}" x2="{ix(xt)}" y2="{y+8}"/>')
    by = y + 14
    parts += [
        f'<line class="thin" x1="{ix(X0)}" y1="{by}" x2="{ix(X0)}" y2="{by+10}"/>',
        f'<line class="thin" x1="{ix(X0)}" y1="{by+10}" x2="{ix(X_CA_END)}" y2="{by+10}"/>',
        f'<line class="thin" x1="{ix(X_CA_END)}" y1="{by}" x2="{ix(X_CA_END)}" y2="{by+10}"/>',
        f'<text x="{ix((X0+X_CA_END)/2)}" y="{by+26}" class="t" text-anchor="middle">(CA)</text>',
        f'<circle cx="{ix(X_BMP)}" cy="{y}" r="{r}" fill="#000"/>',
        f'<text x="{ix(X_BMP)}" y="{y-16}" class="h" text-anchor="middle">BMP</text>',
        f'<line class="thin" x1="{ix(X_2S)}" y1="{y-11}" x2="{ix(X_2S)}" y2="{y+11}"/>',
        f'<text x="{ix(X_2S)}" y="{y+26}" class="t" text-anchor="middle">2S</text>',
        f'<text x="{ix((X_2S+X_INDUC_END)/2)}" y="{y-18}" class="h" text-anchor="middle">BMP + SCF + EGF</text>',
    ]

    for xd in DOTS_MAIN:
        parts.append(f'<circle cx="{ix(xd)}" cy="{y}" r="{r}" fill="#000"/>')

    for xd, yb_in in zip(DOTS_MAIN[:3], BRANCH_YS):
        yb = iy(yb_in)
        s = X_MARK * 100 / 2
        parts += [
            f'<line class="thin dash" x1="{ix(xd)}" y1="{y+r+2}" x2="{ix(xd)}" y2="{yb}"/>',
            f'<line class="med" x1="{ix(xd)}" y1="{yb}" x2="{ix(xd+BRANCH_LEN)}" y2="{yb}"/>',
            f'<line class="med" x1="{ix(xd)-s}" y1="{yb-s}" x2="{ix(xd)+s}" y2="{yb+s}"/>',
            f'<line class="med" x1="{ix(xd)-s}" y1="{yb+s}" x2="{ix(xd)+s}" y2="{yb-s}"/>',
            f'<circle cx="{ix(xd+BRANCH_LEN)}" cy="{yb}" r="{r}" fill="#000"/>',
        ]

    lx, ly = ix(0.55), iy(3.70)
    parts += [
        f'<line class="med" x1="{lx-4}" y1="{ly-4}" x2="{lx+4}" y2="{ly+4}"/>',
        f'<line class="med" x1="{lx-4}" y1="{ly+4}" x2="{lx+4}" y2="{ly-4}"/>',
        f'<text x="{lx+12}" y="{ly+4}" class="s">Ecto (SB + LDN)  ·  Meso  ·  Endo</text>',
        f'<circle cx="{lx}" cy="{ly+24}" r="4.2" fill="#000"/>',
        f'<text x="{lx+12}" y="{ly+28}" class="s">fix &amp; stain</text>',
        "</svg>",
    ]
    path.write_text("\n".join(parts), encoding="utf-8")
    print(f"Wrote {path}")


def build_png(path: Path, scale: int = 3) -> None:
    W, H = int(SLIDE_W * 100 * scale), int(SLIDE_H * 100 * scale)
    im = Image.new("RGB", (W, H), "white")
    dr = ImageDraw.Draw(im)

    def font(size, bold=False):
        candidates = (
            ["/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"]
            if bold
            else ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]
        ) + ["/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"]
        for n in candidates:
            try:
                return ImageFont.truetype(n, size)
            except OSError:
                continue
        return ImageFont.load_default()

    fb = font(14 * scale, bold=True)
    fh = font(12 * scale, bold=True)
    ft = font(11 * scale)
    fs = font(10 * scale)

    def px(v):
        return v * 100 * scale

    def line(x1, y1, x2, y2, w=1.5):
        dr.line([(px(x1), px(y1)), (px(x2), px(y2))], fill="black", width=max(1, int(w * scale)))

    def dash_v(x, y1, y2, w=1.15):
        yy, y_end = px(y1), px(y2)
        step, gap = 7 * scale, 5 * scale
        while yy < y_end:
            y2s = min(yy + step, y_end)
            dr.line([(px(x), yy), (px(x), y2s)], fill="black", width=max(1, int(w * scale)))
            yy = y2s + gap

    def dot(x, y, r=DOT_R):
        cx, cy, rr = px(x), px(y), r * 100 * scale
        dr.ellipse([cx - rr, cy - rr, cx + rr, cy + rr], fill="black")

    def xmark(x, y, size=X_MARK):
        s = size / 2
        line(x - s, y - s, x + s, y + s, w=1.55)
        line(x - s, y + s, x + s, y - s, w=1.55)

    def centered(text, x, y, fnt, fill="black"):
        tw = dr.textlength(text, font=fnt)
        dr.text((px(x) - tw / 2, px(y)), text, font=fnt, fill=fill)

    dr.text((px(0.15), px(0.10)), "A", font=fb, fill="black")
    dr.text((px(0.42), px(0.28)), "PGC induction:", font=fb, fill="black")
    centered("2d FA", 1.40, 0.00, ft)
    line(1.40, 0.20, 1.40, 0.35, w=1.15)
    ax, ay = px(1.40), px(0.42)
    dr.polygon([(ax, ay), (ax - 4.5 * scale, ay - 8 * scale), (ax + 4.5 * scale, ay - 8 * scale)], fill="black")

    line(X0, Y_MAIN, X_INDUC_END, Y_MAIN, w=1.55)
    for xt in (X0, X_CA_END):
        line(xt, Y_MAIN - 0.08, xt, Y_MAIN + 0.08, w=1.15)
    line(X0, Y_MAIN + 0.14, X0, Y_MAIN + 0.24, w=1.15)
    line(X0, Y_MAIN + 0.24, X_CA_END, Y_MAIN + 0.24, w=1.15)
    line(X_CA_END, Y_MAIN + 0.14, X_CA_END, Y_MAIN + 0.24, w=1.15)
    centered("(CA)", (X0 + X_CA_END) / 2, Y_MAIN + 0.28, ft)

    dot(X_BMP, Y_MAIN)
    centered("BMP", X_BMP, Y_MAIN - 0.34, fh)
    line(X_2S, Y_MAIN - 0.11, X_2S, Y_MAIN + 0.11, w=1.15)
    centered("2S", X_2S, Y_MAIN + 0.14, ft)
    centered("BMP + SCF + EGF", (X_2S + X_INDUC_END) / 2, Y_MAIN - 0.40, fh)

    for xd in DOTS_MAIN:
        dot(xd, Y_MAIN)

    for xd, yb in zip(DOTS_MAIN[:3], BRANCH_YS):
        dash_v(xd, Y_MAIN + DOT_R + 0.02, yb)
        line(xd, yb, xd + BRANCH_LEN, yb, w=1.55)
        xmark(xd, yb)
        dot(xd + BRANCH_LEN, yb)

    xmark(0.55, 3.70, size=0.075)
    dr.text((px(0.72), px(3.62)), "Ecto (SB + LDN)  ·  Meso  ·  Endo", font=fs, fill="#333333")
    dot(0.55, 3.95, r=0.042)
    dr.text((px(0.72), px(3.86)), "fix & stain", font=fs, fill="#333333")

    im.save(path, "PNG")
    print(f"Wrote {path}")


def main():
    build_pptx(OUT_DIR / "fig1a_pgc_induction.pptx")
    build_svg(OUT_DIR / "fig1a_pgc_induction.svg")
    build_png(OUT_DIR / "fig1a_pgc_induction.png")


if __name__ == "__main__":
    main()
